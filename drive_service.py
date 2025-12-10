"""
Google Drive Service
Handles fetching and processing documents from Google Drive
"""

import io
import os
from typing import List, Dict, Optional
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from googleapiclient.errors import HttpError

# Document processing
from pypdf import PdfReader
from docx import Document as DocxDocument
from bs4 import BeautifulSoup
import openpyxl


# Supported MIME types for processing
SUPPORTED_MIME_TYPES = {
    # Google Workspace
    'application/vnd.google-apps.document': 'gdoc',
    'application/vnd.google-apps.spreadsheet': 'gsheet',
    'application/vnd.google-apps.presentation': 'gslides',

    # Microsoft Office - Modern formats
    'application/vnd.openxmlformats-officedocument.wordprocessingml.document': 'docx',
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': 'xlsx',
    'application/vnd.openxmlformats-officedocument.presentationml.presentation': 'pptx',

    # Microsoft Office - Legacy formats
    'application/msword': 'doc',
    'application/vnd.ms-excel': 'xls',
    'application/vnd.ms-powerpoint': 'ppt',

    # PDF
    'application/pdf': 'pdf',

    # Text formats
    'text/plain': 'txt',
    'text/html': 'html',
    'text/markdown': 'md',
    'text/csv': 'csv',
    'text/xml': 'xml',
    'text/rtf': 'rtf',
    'application/rtf': 'rtf',

    # Data formats
    'application/json': 'json',
    'application/xml': 'xml',
    'text/x-python': 'py',
    'text/x-java-source': 'java',
    'text/javascript': 'js',
    'application/javascript': 'js',
    'text/css': 'css',
    'text/x-c': 'c',
    'text/x-c++src': 'cpp',

    # Other common formats
    'application/x-yaml': 'yaml',
    'text/yaml': 'yaml',
    'text/x-rst': 'rst',
    'text/x-log': 'log',
}

# Export formats for Google Workspace documents
EXPORT_MIME_TYPES = {
    'application/vnd.google-apps.document': 'text/plain',
    'application/vnd.google-apps.spreadsheet': 'text/csv',
    'application/vnd.google-apps.presentation': 'text/plain',
}


class DriveService:
    """Service for interacting with Google Drive API"""

    def __init__(self, credentials):
        """Initialize Drive service with credentials"""
        self.credentials = credentials
        # Initialize all Google services
        self.service = build('drive', 'v3', credentials=credentials)  # Drive API
        self.docs_service = build('docs', 'v1', credentials=credentials)  # Docs API
        self.sheets_service = build('sheets', 'v4', credentials=credentials)  # Sheets API
        self.slides_service = build('slides', 'v1', credentials=credentials)  # Slides API

    def list_folders(self) -> List[Dict]:
        """
        List all folders from Google Drive

        Returns:
            List of folder metadata dictionaries
        """
        all_folders = []
        page_token = None

        try:
            while True:
                results = self.service.files().list(
                    pageSize=100,
                    fields="nextPageToken, files(id, name, mimeType)",
                    q="mimeType='application/vnd.google-apps.folder' and trashed=false",
                    pageToken=page_token
                ).execute()

                folders = results.get('files', [])
                all_folders.extend(folders)

                page_token = results.get('nextPageToken')
                if not page_token:
                    break

            return all_folders

        except HttpError as error:
            print(f"Error listing folders: {error}")
            return []

    def list_folder_contents(self, folder_id: Optional[str] = None) -> List[Dict]:
        """
        List contents of a folder (files and subfolders)

        Args:
            folder_id: Folder ID to list contents from. If None, lists root.

        Returns:
            List of file/folder metadata with type indicator
        """
        all_items = []
        page_token = None

        if folder_id:
            query = f"'{folder_id}' in parents and trashed=false"
        else:
            query = "'root' in parents and trashed=false"

        try:
            while True:
                results = self.service.files().list(
                    pageSize=100,
                    fields="nextPageToken, files(id, name, mimeType, modifiedTime, size)",
                    q=query,
                    orderBy="folder,name",
                    pageToken=page_token
                ).execute()

                items = results.get('files', [])
                for item in items:
                    is_folder = item['mimeType'] == 'application/vnd.google-apps.folder'
                    is_supported = item['mimeType'] in SUPPORTED_MIME_TYPES

                    # Show all files (not just supported ones) - we'll try to index them anyway
                    all_items.append({
                        'id': item['id'],
                        'name': item['name'],
                        'mimeType': item['mimeType'],
                        'type': 'folder' if is_folder else 'file',
                        'modifiedTime': item.get('modifiedTime', ''),
                        'size': item.get('size', 0),
                        'supported': is_folder or is_supported  # Flag to show in UI
                    })

                page_token = results.get('nextPageToken')
                if not page_token:
                    break

            return all_items

        except HttpError as error:
            print(f"Error listing folder contents: {error}")
            return []

    def get_documents_by_ids(self, file_ids: List[str]) -> List[Dict]:
        """
        Fetch specific documents by their IDs

        Args:
            file_ids: List of file IDs to fetch

        Returns:
            List of documents with metadata and content
        """
        documents = []

        for file_id in file_ids:
            try:
                # Get file metadata
                print(f"[get_documents_by_ids] Fetching metadata for {file_id}...")
                file_meta = self.service.files().get(
                    fileId=file_id,
                    fields="id, name, mimeType, modifiedTime"
                ).execute()

                mime_type = file_meta['mimeType']
                file_name = file_meta['name']
                print(f"[get_documents_by_ids] File: {file_name} | Type: {mime_type}")

                # Skip folders
                if mime_type == 'application/vnd.google-apps.folder':
                    print(f"[get_documents_by_ids] Skipping folder")
                    continue

                # Try to process any file type (no longer skip unsupported types)
                print(f"[get_documents_by_ids] Getting content...")
                content = self.get_file_content(file_id, mime_type)
                print(f"[get_documents_by_ids] Content length: {len(content) if content else 0}")

                if content and content.strip():
                    documents.append({
                        'id': file_id,
                        'name': file_name,
                        'mime_type': mime_type,
                        'content': content,
                        'modified_time': file_meta.get('modifiedTime', '')
                    })
                    print(f"[get_documents_by_ids] Document added successfully")

            except HttpError as error:
                print(f"[get_documents_by_ids] HttpError for {file_id}: {error}")
                continue
            except Exception as e:
                import traceback
                print(f"[get_documents_by_ids] ERROR for {file_id}: {e}\n{traceback.format_exc()}")
                continue

        return documents

    def list_files(self, page_size: int = 100, folder_id: Optional[str] = None) -> List[Dict]:
        """
        List files from Google Drive

        Args:
            page_size: Number of files to fetch per page
            folder_id: Optional folder ID to list files from

        Returns:
            List of file metadata dictionaries
        """
        all_files = []
        page_token = None

        # List all files (not filtered by MIME type) to support any format
        if folder_id:
            query = f"'{folder_id}' in parents and trashed=false"
        else:
            query = "trashed=false"

        try:
            while True:
                results = self.service.files().list(
                    pageSize=page_size,
                    fields="nextPageToken, files(id, name, mimeType, modifiedTime, size, parents)",
                    q=query,
                    pageToken=page_token
                ).execute()

                files = results.get('files', [])
                all_files.extend(files)

                page_token = results.get('nextPageToken')
                if not page_token:
                    break

            return all_files

        except HttpError as error:
            print(f"Error listing files: {error}")
            return []

    def list_files_recursive(self, folder_id: str) -> List[Dict]:
        """
        Recursively list all files in a folder and its subfolders

        Args:
            folder_id: The folder ID to start from

        Returns:
            List of file metadata dictionaries
        """
        all_files = []

        def process_folder(fid: str):
            items = self.list_files(folder_id=fid)
            for item in items:
                if item['mimeType'] == 'application/vnd.google-apps.folder':
                    # Recursively process subfolder
                    process_folder(item['id'])
                else:
                    # Include all files (not just supported types)
                    all_files.append(item)

        process_folder(folder_id)
        return all_files

    def get_file_content(self, file_id: str, mime_type: str) -> Optional[str]:
        """
        Get content of a file from Google Drive

        Args:
            file_id: The ID of the file to fetch
            mime_type: The MIME type of the file

        Returns:
            Text content of the file or None if failed
        """
        try:
            print(f"[get_file_content] Starting for {file_id}, type: {mime_type}")
            # For Google Workspace documents, export as text
            if mime_type in EXPORT_MIME_TYPES:
                print(f"[get_file_content] Using export method (Google Workspace doc)")
                content = self._export_google_doc(file_id, mime_type)
                print(f"[get_file_content] Export done, content length: {len(content) if content else 0}")
                return content
            else:
                print(f"[get_file_content] Using download method")
                content = self._download_and_parse(file_id, mime_type)
                print(f"[get_file_content] Download done, content length: {len(content) if content else 0}")
                return content

        except HttpError as error:
            import traceback
            print(f"[get_file_content] HttpError for {file_id}: {error}\n{traceback.format_exc()}")
            return None
        except Exception as e:
            import traceback
            print(f"[get_file_content] ERROR for {file_id}: {e}\n{traceback.format_exc()}")
            return None

    def _export_google_doc(self, file_id: str, mime_type: str) -> Optional[str]:
        """Export Google Workspace document to text"""
        try:
            print(f"[_export_google_doc] Starting export for {file_id}")
            export_mime = EXPORT_MIME_TYPES.get(mime_type, 'text/plain')
            print(f"[_export_google_doc] Export MIME type: {export_mime}")

            request = self.service.files().export_media(
                fileId=file_id,
                mimeType=export_mime
            )

            file_content = io.BytesIO()
            downloader = MediaIoBaseDownload(file_content, request)

            print(f"[_export_google_doc] Starting download chunks...")
            done = False
            chunk_num = 0
            while not done:
                status, done = downloader.next_chunk()
                chunk_num += 1
                if status:
                    print(f"[_export_google_doc] Chunk {chunk_num}: {int(status.progress() * 100)}%")

            print(f"[_export_google_doc] Download complete, decoding...")
            content = file_content.getvalue().decode('utf-8', errors='ignore')
            print(f"[_export_google_doc] Done! Content length: {len(content)}")
            return content

        except Exception as e:
            import traceback
            print(f"[_export_google_doc] ERROR: {e}\n{traceback.format_exc()}")
            return ""

    def _download_and_parse(self, file_id: str, mime_type: str) -> Optional[str]:
        """Download and parse a file based on its MIME type"""
        request = self.service.files().get_media(fileId=file_id)

        file_content = io.BytesIO()
        downloader = MediaIoBaseDownload(file_content, request)

        done = False
        while not done:
            status, done = downloader.next_chunk()

        file_content.seek(0)

        # Parse based on MIME type
        if mime_type == 'application/pdf':
            return self._parse_pdf(file_content)
        elif mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
            return self._parse_docx(file_content)
        elif mime_type in ['application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'application/vnd.ms-excel']:
            return self._parse_xlsx(file_content)
        elif mime_type == 'application/vnd.openxmlformats-officedocument.presentationml.presentation':
            return self._parse_pptx(file_content)
        elif mime_type == 'text/html':
            return self._parse_html(file_content)
        elif mime_type == 'text/csv':
            return self._parse_csv(file_content)
        elif mime_type in ['application/json']:
            return self._parse_json(file_content)
        elif mime_type in ['application/xml', 'text/xml']:
            return self._parse_xml(file_content)
        else:
            # Fallback: try to read as text
            return self._parse_as_text(file_content)

    def _parse_pdf(self, file_content: io.BytesIO) -> str:
        """Extract text from PDF file"""
        try:
            reader = PdfReader(file_content)
            text_parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            return "\n\n".join(text_parts)
        except Exception as e:
            print(f"Error parsing PDF: {e}")
            return ""

    def _parse_docx(self, file_content: io.BytesIO) -> str:
        """Extract text from DOCX file"""
        try:
            print("  -> DOCX: Loading document...")
            doc = DocxDocument(file_content)
            print(f"  -> DOCX: Found {len(doc.paragraphs)} paragraphs")
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            result = "\n\n".join(text_parts)
            print(f"  -> DOCX: Extracted {len(result)} characters")
            return result
        except Exception as e:
            import traceback
            print(f"ERROR parsing DOCX: {e}\n{traceback.format_exc()}")
            return ""

    def _parse_xlsx(self, file_content: io.BytesIO) -> str:
        """Extract text from XLSX file"""
        try:
            workbook = openpyxl.load_workbook(file_content, read_only=True)
            text_parts = []

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text_parts.append(f"=== Sheet: {sheet_name} ===")

                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
                    if row_text.strip():
                        text_parts.append(row_text)

            return "\n".join(text_parts)
        except Exception as e:
            print(f"Error parsing XLSX: {e}")
            return ""

    def _parse_html(self, file_content: io.BytesIO) -> str:
        """Extract text from HTML file"""
        try:
            soup = BeautifulSoup(file_content.read(), 'html.parser')
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
            text = soup.get_text(separator='\n')
            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)
            return text
        except Exception as e:
            print(f"Error parsing HTML: {e}")
            return ""

    def _parse_pptx(self, file_content: io.BytesIO) -> str:
        """Extract text from PPTX file"""
        try:
            from pptx import Presentation
            prs = Presentation(file_content)
            text_parts = []

            for slide_num, slide in enumerate(prs.slides, 1):
                text_parts.append(f"=== Slide {slide_num} ===")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text)

            return "\n\n".join(text_parts)
        except ImportError:
            print("python-pptx not installed, trying fallback")
            return self._parse_as_text(file_content)
        except Exception as e:
            print(f"Error parsing PPTX: {e}")
            return ""

    def _parse_csv(self, file_content: io.BytesIO) -> str:
        """Extract text from CSV file"""
        try:
            import csv
            content = file_content.read().decode('utf-8', errors='ignore')
            file_content.seek(0)

            text_parts = []
            reader = csv.reader(content.splitlines())
            for row in reader:
                row_text = " | ".join(row)
                if row_text.strip():
                    text_parts.append(row_text)

            return "\n".join(text_parts)
        except Exception as e:
            print(f"Error parsing CSV: {e}")
            return file_content.read().decode('utf-8', errors='ignore')

    def _parse_json(self, file_content: io.BytesIO) -> str:
        """Extract text from JSON file"""
        try:
            import json
            content = file_content.read().decode('utf-8', errors='ignore')
            data = json.loads(content)
            # Pretty print JSON for better readability
            return json.dumps(data, indent=2, ensure_ascii=False)
        except Exception as e:
            print(f"Error parsing JSON: {e}")
            file_content.seek(0)
            return file_content.read().decode('utf-8', errors='ignore')

    def _parse_xml(self, file_content: io.BytesIO) -> str:
        """Extract text from XML file"""
        try:
            soup = BeautifulSoup(file_content.read(), 'xml')
            text = soup.get_text(separator='\n')
            lines = (line.strip() for line in text.splitlines())
            return '\n'.join(line for line in lines if line)
        except Exception as e:
            print(f"Error parsing XML: {e}")
            file_content.seek(0)
            return file_content.read().decode('utf-8', errors='ignore')

    def _parse_as_text(self, file_content: io.BytesIO) -> str:
        """Fallback: try to read any file as text"""
        try:
            content = file_content.read()
            # Try UTF-8 first
            try:
                return content.decode('utf-8')
            except UnicodeDecodeError:
                pass
            # Try Latin-1 (covers all byte values)
            try:
                return content.decode('latin-1')
            except UnicodeDecodeError:
                pass
            # Last resort: ignore errors
            return content.decode('utf-8', errors='ignore')
        except Exception as e:
            print(f"Error reading file as text: {e}")
            return ""

    def get_all_documents(self, folder_id: Optional[str] = None) -> List[Dict]:
        """
        Fetch all documents and their content from Drive

        Args:
            folder_id: Optional folder ID to fetch documents from (recursive)

        Returns:
            List of documents with metadata and content
        """
        if folder_id:
            files = self.list_files_recursive(folder_id)
        else:
            files = self.list_files()

        documents = []

        for file in files:
            file_id = file['id']
            file_name = file['name']
            mime_type = file['mimeType']

            print(f"Processing: {file_name}")

            content = self.get_file_content(file_id, mime_type)

            if content and content.strip():
                documents.append({
                    'id': file_id,
                    'name': file_name,
                    'mime_type': mime_type,
                    'content': content,
                    'modified_time': file.get('modifiedTime', '')
                })

        return documents

    def get_user_info(self) -> Dict:
        """Get information about the authenticated user's Drive"""
        try:
            about = self.service.about().get(fields="user, storageQuota").execute()
            return about
        except HttpError as error:
            print(f"Error getting user info: {error}")
            return {}

    def get_document_content(self, file_id: str) -> str:
        """
        Get the text content of a Google Doc

        Args:
            file_id: The ID of the Google Doc

        Returns:
            The text content of the document
        """
        try:
            doc = self.docs_service.documents().get(documentId=file_id).execute()

            # Extract text from all elements
            text_parts = []
            for element in doc.get('body', {}).get('content', []):
                if 'paragraph' in element:
                    for elem in element['paragraph'].get('elements', []):
                        if 'textRun' in elem:
                            text_parts.append(elem['textRun'].get('content', ''))

            return ''.join(text_parts)

        except Exception as e:
            print(f"[get_document_content] Error: {e}", flush=True)
            return ""

    def update_google_doc(self, file_id: str, new_content: str) -> Dict:
        """
        Update a Google Doc with new content (replaces entire content)

        Args:
            file_id: The ID of the Google Doc
            new_content: The new text content

        Returns:
            Result dictionary with status
        """
        try:
            # First, get the document to find its content length
            doc = self.docs_service.documents().get(documentId=file_id).execute()

            # Get the end index of the document content
            content = doc.get('body', {}).get('content', [])
            end_index = 1
            for element in content:
                if 'endIndex' in element:
                    end_index = max(end_index, element['endIndex'])

            # Prepare requests to clear and insert new content
            requests = []

            # Delete existing content (if any beyond the initial newline)
            if end_index > 2:
                requests.append({
                    'deleteContentRange': {
                        'range': {
                            'startIndex': 1,
                            'endIndex': end_index - 1
                        }
                    }
                })

            # Insert new content at the beginning
            if new_content:
                requests.append({
                    'insertText': {
                        'location': {
                            'index': 1
                        },
                        'text': new_content
                    }
                })

            # Execute the batch update
            if requests:
                self.docs_service.documents().batchUpdate(
                    documentId=file_id,
                    body={'requests': requests}
                ).execute()

            print(f"[update_google_doc] Successfully updated document {file_id}", flush=True)
            return {
                'status': 'success',
                'message': 'Document updated successfully',
                'file_id': file_id
            }

        except HttpError as error:
            print(f"[update_google_doc] HttpError: {error}", flush=True)
            return {
                'status': 'error',
                'message': str(error),
                'file_id': file_id
            }
        except Exception as e:
            import traceback
            print(f"[update_google_doc] Error: {e}\n{traceback.format_exc()}", flush=True)
            return {
                'status': 'error',
                'message': str(e),
                'file_id': file_id
            }

    def append_to_google_doc(self, file_id: str, content_to_append: str) -> Dict:
        """
        Append content to the end of a Google Doc

        Args:
            file_id: The ID of the Google Doc
            content_to_append: The text to append

        Returns:
            Result dictionary with status
        """
        try:
            # Get current document end index
            doc = self.docs_service.documents().get(documentId=file_id).execute()
            content = doc.get('body', {}).get('content', [])
            end_index = 1
            for element in content:
                if 'endIndex' in element:
                    end_index = max(end_index, element['endIndex'])

            # Insert at the end (before the final newline)
            insert_index = max(1, end_index - 1)

            requests = [{
                'insertText': {
                    'location': {
                        'index': insert_index
                    },
                    'text': '\n' + content_to_append
                }
            }]

            self.docs_service.documents().batchUpdate(
                documentId=file_id,
                body={'requests': requests}
            ).execute()

            print(f"[append_to_google_doc] Successfully appended to document {file_id}", flush=True)
            return {
                'status': 'success',
                'message': 'Content appended successfully',
                'file_id': file_id
            }

        except Exception as e:
            import traceback
            print(f"[append_to_google_doc] Error: {e}\n{traceback.format_exc()}", flush=True)
            return {
                'status': 'error',
                'message': str(e),
                'file_id': file_id
            }

    def replace_text_in_doc(self, file_id: str, find_text: str, replace_text: str) -> Dict:
        """
        Find and replace specific text in a Google Doc

        Args:
            file_id: The ID of the Google Doc
            find_text: The text to find
            replace_text: The text to replace it with

        Returns:
            Result dictionary with status
        """
        try:
            requests = [{
                'replaceAllText': {
                    'containsText': {
                        'text': find_text,
                        'matchCase': False
                    },
                    'replaceText': replace_text
                }
            }]

            result = self.docs_service.documents().batchUpdate(
                documentId=file_id,
                body={'requests': requests}
            ).execute()

            # Check how many replacements were made
            replies = result.get('replies', [])
            occurrences = 0
            if replies and 'replaceAllText' in replies[0]:
                occurrences = replies[0]['replaceAllText'].get('occurrencesChanged', 0)

            print(f"[replace_text_in_doc] Replaced {occurrences} occurrences in {file_id}", flush=True)
            return {
                'status': 'success',
                'message': f'Replaced {occurrences} occurrence(s) of "{find_text}"',
                'occurrences': occurrences,
                'file_id': file_id
            }

        except Exception as e:
            import traceback
            print(f"[replace_text_in_doc] Error: {e}\n{traceback.format_exc()}", flush=True)
            return {
                'status': 'error',
                'message': str(e),
                'file_id': file_id
            }

    def insert_text_after(self, file_id: str, after_text: str, new_text: str) -> Dict:
        """
        Insert text after a specific text in a Google Doc

        Args:
            file_id: The ID of the Google Doc
            after_text: The text after which to insert
            new_text: The text to insert

        Returns:
            Result dictionary with status
        """
        try:
            # Get the document content
            doc = self.docs_service.documents().get(documentId=file_id).execute()

            # Find the position of after_text
            full_text = ""
            for element in doc.get('body', {}).get('content', []):
                if 'paragraph' in element:
                    for elem in element['paragraph'].get('elements', []):
                        if 'textRun' in elem:
                            full_text += elem['textRun'].get('content', '')

            # Find the index where to insert
            position = full_text.find(after_text)
            if position == -1:
                return {
                    'status': 'error',
                    'message': f'Text "{after_text}" not found in document',
                    'file_id': file_id
                }

            # Calculate the insertion index (+1 for document offset, + length of found text)
            insert_index = position + len(after_text) + 1

            requests = [{
                'insertText': {
                    'location': {
                        'index': insert_index
                    },
                    'text': new_text
                }
            }]

            self.docs_service.documents().batchUpdate(
                documentId=file_id,
                body={'requests': requests}
            ).execute()

            print(f"[insert_text_after] Inserted text after '{after_text}' in {file_id}", flush=True)
            return {
                'status': 'success',
                'message': f'Text inserted after "{after_text}"',
                'file_id': file_id
            }

        except Exception as e:
            import traceback
            print(f"[insert_text_after] Error: {e}\n{traceback.format_exc()}", flush=True)
            return {
                'status': 'error',
                'message': str(e),
                'file_id': file_id
            }

    def delete_text_in_doc(self, file_id: str, text_to_delete: str) -> Dict:
        """
        Delete specific text from a Google Doc

        Args:
            file_id: The ID of the Google Doc
            text_to_delete: The text to delete

        Returns:
            Result dictionary with status
        """
        try:
            # Use replace with empty string to delete
            return self.replace_text_in_doc(file_id, text_to_delete, '')

        except Exception as e:
            import traceback
            print(f"[delete_text_in_doc] Error: {e}\n{traceback.format_exc()}", flush=True)
            return {
                'status': 'error',
                'message': str(e),
                'file_id': file_id
            }

    def create_google_doc(self, title: str, content: str = "") -> Dict:
        """
        Create a new Google Doc

        Args:
            title: The title of the new document
            content: Optional initial content

        Returns:
            Result dictionary with file_id and status
        """
        try:
            # Create empty document
            doc = self.docs_service.documents().create(body={'title': title}).execute()
            file_id = doc.get('documentId')

            # Add content if provided
            if content:
                requests = [{
                    'insertText': {
                        'location': {'index': 1},
                        'text': content
                    }
                }]
                self.docs_service.documents().batchUpdate(
                    documentId=file_id,
                    body={'requests': requests}
                ).execute()

            print(f"[create_google_doc] Created document: {file_id}", flush=True)
            return {
                'status': 'success',
                'message': f'Document "{title}" created successfully',
                'file_id': file_id,
                'title': title
            }

        except Exception as e:
            import traceback
            print(f"[create_google_doc] Error: {e}\n{traceback.format_exc()}", flush=True)
            return {
                'status': 'error',
                'message': str(e)
            }

    def get_file_metadata(self, file_id: str) -> Dict:
        """
        Get metadata for a file

        Args:
            file_id: The file ID

        Returns:
            File metadata dictionary
        """
        try:
            file_meta = self.service.files().get(
                fileId=file_id,
                fields="id, name, mimeType, modifiedTime, webViewLink"
            ).execute()
            return file_meta
        except HttpError as error:
            print(f"[get_file_metadata] Error: {error}", flush=True)
            return {}

    # ============== Google Sheets Methods ==============

    def update_google_sheet(self, file_id: str, data: list, sheet_name: str = None, clear_first: bool = False) -> Dict:
        """
        Update a Google Sheet with data

        Args:
            file_id: The ID of the Google Sheet
            data: 2D list of values [[row1], [row2], ...]
            sheet_name: Optional sheet name (defaults to first sheet)
            clear_first: If True, clear the sheet before adding data

        Returns:
            Result dictionary with status
        """
        try:
            # Default range
            range_name = f"'{sheet_name}'!A1" if sheet_name else "A1"

            if clear_first:
                # Clear the sheet first
                clear_range = f"'{sheet_name}'" if sheet_name else "A:ZZ"
                self.sheets_service.spreadsheets().values().clear(
                    spreadsheetId=file_id,
                    range=clear_range,
                    body={}
                ).execute()

            # Update with new data
            body = {'values': data}
            result = self.sheets_service.spreadsheets().values().update(
                spreadsheetId=file_id,
                range=range_name,
                valueInputOption='USER_ENTERED',
                body=body
            ).execute()

            print(f"[update_google_sheet] Updated {result.get('updatedCells', 0)} cells", flush=True)
            return {
                'status': 'success',
                'message': f"Updated {result.get('updatedCells', 0)} cells",
                'file_id': file_id
            }

        except Exception as e:
            import traceback
            print(f"[update_google_sheet] Error: {e}\n{traceback.format_exc()}", flush=True)
            return {'status': 'error', 'message': str(e), 'file_id': file_id}

    def append_to_google_sheet(self, file_id: str, data: list, sheet_name: str = None) -> Dict:
        """
        Append rows to a Google Sheet

        Args:
            file_id: The ID of the Google Sheet
            data: 2D list of values to append [[row1], [row2], ...]
            sheet_name: Optional sheet name

        Returns:
            Result dictionary with status
        """
        try:
            range_name = f"'{sheet_name}'!A:A" if sheet_name else "A:A"

            body = {'values': data}
            result = self.sheets_service.spreadsheets().values().append(
                spreadsheetId=file_id,
                range=range_name,
                valueInputOption='USER_ENTERED',
                insertDataOption='INSERT_ROWS',
                body=body
            ).execute()

            print(f"[append_to_google_sheet] Appended rows", flush=True)
            return {
                'status': 'success',
                'message': f"Appended {len(data)} rows",
                'file_id': file_id
            }

        except Exception as e:
            import traceback
            print(f"[append_to_google_sheet] Error: {e}\n{traceback.format_exc()}", flush=True)
            return {'status': 'error', 'message': str(e), 'file_id': file_id}

    def get_sheet_content(self, file_id: str, sheet_name: str = None) -> Dict:
        """
        Get content from a Google Sheet

        Args:
            file_id: The ID of the Google Sheet
            sheet_name: Optional sheet name

        Returns:
            Dict with values and status
        """
        try:
            range_name = f"'{sheet_name}'" if sheet_name else "A:ZZ"

            result = self.sheets_service.spreadsheets().values().get(
                spreadsheetId=file_id,
                range=range_name
            ).execute()

            values = result.get('values', [])
            return {
                'status': 'success',
                'values': values,
                'rows': len(values)
            }

        except Exception as e:
            print(f"[get_sheet_content] Error: {e}", flush=True)
            return {'status': 'error', 'message': str(e), 'values': []}

    def update_sheet_cells(self, file_id: str, updates: list) -> Dict:
        """
        Update specific cells in a Google Sheet

        Args:
            file_id: The ID of the Google Sheet
            updates: List of cell updates [{"cell": "A1", "value": "new value"}, ...]

        Returns:
            Result dictionary with status
        """
        try:
            # Prepare batch update data
            data = []
            for update in updates:
                cell = update.get('cell', '')
                value = update.get('value', '')
                if cell:
                    data.append({
                        'range': cell,
                        'values': [[value]]
                    })

            if not data:
                return {'status': 'error', 'message': 'No valid updates provided'}

            body = {
                'valueInputOption': 'USER_ENTERED',
                'data': data
            }

            result = self.sheets_service.spreadsheets().values().batchUpdate(
                spreadsheetId=file_id,
                body=body
            ).execute()

            updated_cells = result.get('totalUpdatedCells', 0)
            print(f"[update_sheet_cells] Updated {updated_cells} cells", flush=True)
            return {
                'status': 'success',
                'message': f'{updated_cells} cellule(s) mise(s) a jour',
                'updated_cells': updated_cells
            }

        except Exception as e:
            import traceback
            print(f"[update_sheet_cells] Error: {e}\n{traceback.format_exc()}", flush=True)
            return {'status': 'error', 'message': str(e)}

    # ============== Google Slides Methods ==============

    def update_google_slides(self, file_id: str, updates: list) -> Dict:
        """
        Update a Google Slides presentation

        Args:
            file_id: The ID of the Google Slides presentation
            updates: List of update requests

        Returns:
            Result dictionary with status
        """
        try:
            body = {'requests': updates}
            result = self.slides_service.presentations().batchUpdate(
                presentationId=file_id,
                body=body
            ).execute()

            print(f"[update_google_slides] Updated presentation", flush=True)
            return {
                'status': 'success',
                'message': 'Presentation updated',
                'file_id': file_id
            }

        except Exception as e:
            import traceback
            print(f"[update_google_slides] Error: {e}\n{traceback.format_exc()}", flush=True)
            return {'status': 'error', 'message': str(e), 'file_id': file_id}

    def add_slide_with_text(self, file_id: str, title: str, body_text: str) -> Dict:
        """
        Add a new slide with title and body text

        Args:
            file_id: The ID of the Google Slides presentation
            title: Slide title
            body_text: Slide body content

        Returns:
            Result dictionary with status
        """
        try:
            # Create a new slide
            requests = [
                {
                    'createSlide': {
                        'slideLayoutReference': {
                            'predefinedLayout': 'TITLE_AND_BODY'
                        }
                    }
                }
            ]

            result = self.slides_service.presentations().batchUpdate(
                presentationId=file_id,
                body={'requests': requests}
            ).execute()

            # Get the new slide ID
            slide_id = result['replies'][0]['createSlide']['objectId']

            # Get the presentation to find placeholder IDs
            presentation = self.slides_service.presentations().get(
                presentationId=file_id
            ).execute()

            # Find the new slide and its placeholders
            title_id = None
            body_id = None
            for slide in presentation.get('slides', []):
                if slide['objectId'] == slide_id:
                    for element in slide.get('pageElements', []):
                        if 'shape' in element:
                            placeholder = element['shape'].get('placeholder', {})
                            if placeholder.get('type') == 'TITLE':
                                title_id = element['objectId']
                            elif placeholder.get('type') == 'BODY':
                                body_id = element['objectId']

            # Add text to placeholders
            text_requests = []
            if title_id:
                text_requests.append({
                    'insertText': {
                        'objectId': title_id,
                        'text': title
                    }
                })
            if body_id:
                text_requests.append({
                    'insertText': {
                        'objectId': body_id,
                        'text': body_text
                    }
                })

            if text_requests:
                self.slides_service.presentations().batchUpdate(
                    presentationId=file_id,
                    body={'requests': text_requests}
                ).execute()

            print(f"[add_slide_with_text] Added new slide", flush=True)
            return {
                'status': 'success',
                'message': 'New slide added',
                'file_id': file_id,
                'slide_id': slide_id
            }

        except Exception as e:
            import traceback
            print(f"[add_slide_with_text] Error: {e}\n{traceback.format_exc()}", flush=True)
            return {'status': 'error', 'message': str(e), 'file_id': file_id}
